"""
build_presentation.py - 在已有 15 页 PPT 上增补 3 页并同步数字（W4 交付物，任务 6.4）
项目：信贷违约预测模型

为什么用"克隆 + 改字"而不是重做：
  现有 15 页是成套设计（配色/字号/版式统一）。重画容易走形，克隆母版页
  （第 11 页"泛化能力"：小标题 + 导语 + 表格 + 图片 + 要点框）再替换内容，
  能保证新页和原页风格完全一致。

增补的 3 页（克隆母版 → 替换文字/表格/图片）：
  ① 特征数量实验（W4）：插在"06 特征工程"之后   → 回答"38 个特征是否都用得上"
  ② 概率校准（W4）    ：插在"12 模型解释"之后   → 校准是期望损失框架的前提
  ③ 期望损失框架（W4）：插在"风险阈值选择"之后   → 从一刀切到逐笔算账

同时同步已过时的数字：
  - 导览页"演示最优阈值 0.65" → 改成校准 + 期望损失框架的口径
  - 阈值页"金额账"演示参数（回收率 50%/息差 5%）→ 与 W4 框架统一为 30%/6%
  - 业务建议页"改进方向"里的"用真实客群校准"（已做）→ 改为"迁移到真实客群"

用法：/usr/local/bin/python3 scripts/W4/build_presentation.py
产出：outputs/W4/presentation.pptx（18 页，原地更新）
"""
import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

BASE = next(
    (p for p in [Path(__file__).resolve().parent, *Path(__file__).resolve().parents]
     if (p / "outputs").is_dir()),
    Path(__file__).resolve().parents[2],
)
PPTX = BASE / "outputs" / "W4" / "presentation.pptx"
FIG = BASE / "outputs" / "W4"
RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

prs = Presentation(str(PPTX))
NEW_SLIDES = 3
TOTAL = len(prs.slides) + NEW_SLIDES
print(f"读取: {PPTX.name}（当前 {len(prs.slides)} 页，目标 {TOTAL} 页）")


# ---------- 工具函数 ----------
def text_shapes(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def find_shape(slide, startswith):
    for sh in text_shapes(slide):
        if sh.text_frame.text.strip().startswith(startswith):
            return sh
    raise KeyError(f"未找到以 {startswith!r} 开头的文本框")


def set_para_text(para, text):
    """替换段落文字，保留第一个 run 的字体格式"""
    runs = para.runs
    if not runs:
        para.add_run().text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_lines(shape, lines):
    """按行数写文本：不够就克隆末段，多了就删"""
    tf = shape.text_frame
    proto = copy.deepcopy(tf.paragraphs[-1]._p)
    while len(tf.paragraphs) < len(lines):
        tf._txBody.append(copy.deepcopy(proto))
    while len(tf.paragraphs) > len(lines):
        p = tf.paragraphs[-1]._p
        p.getparent().remove(p)
    for para, line in zip(tf.paragraphs, lines):
        set_para_text(para, line)


def set_cell(cell, text):
    tf = cell.text_frame
    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)
    set_para_text(tf.paragraphs[0], text)


def replace_picture(slide, shape, img_path):
    """换图并保持原始占位框、按比例缩放居中"""
    _, rId = slide.part.get_or_add_image_part(str(img_path))
    shape._element.blipFill.blip.set(RNS, rId)
    img = slide.part.related_part(rId).image
    px_w, px_h = img.size
    ratio = px_h / px_w
    box_l, box_t, box_w, box_h = shape.left, shape.top, shape.width, shape.height
    new_w = box_w
    new_h = int(box_w * ratio)
    if new_h > box_h:
        new_h = box_h
        new_w = int(box_h / ratio)
    shape.width, shape.height = new_w, new_h
    shape.left = box_l + (box_w - new_w) // 2
    shape.top = box_t + (box_h - new_h) // 2


def clone_slide(src_slide):
    """深拷贝幻灯片（含图片关系），返回新页（追加在末尾）"""
    dest = prs.slides.add_slide(src_slide.slide_layout)
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src_slide.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    for rel in src_slide.part.rels.values():
        if rel.is_external:
            dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            dest.part.rels.get_or_add(rel.reltype, rel._target)
    return dest


def reorder(new_order):
    """按给定索引顺序重排幻灯片"""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    for el in ids:
        sldIdLst.remove(el)
    for i in new_order:
        sldIdLst.append(ids[i])


# ---------- 1) 克隆母版（第 11 页 = 泛化能力）并写入新内容 ----------
TEMPLATE = 10  # 0-based，第 11 页
template = prs.slides[TEMPLATE]

CONTENT = [
    dict(tag="特征数量实验（W4）",
         title="38 个特征，真的都用得上吗",
         lead="按随机森林重要性从高到低，每次只喂前 N 个特征重训同一个 LGB，比较验证集 AUC——"
              "只看训练集与验证集，测试集不参与。",
         table=[["用多少特征", "验证集 AUC", "相比满血模型"],
                ["Top 3", "0.7005", "−0.0289（满血的 96%）"],
                ["Top 20", "0.7253", "−0.0041（99.4%）"],
                ["Top 38（最终）", "0.7294", "—"]],
         box=["结论：25 个之后基本饱和",
              "3 个特征（子等级、利率、等级）就拿到满血性能的 96%——机构自己的风险判断最值钱；"
              "20 个到 99.4%；25→30 验证 AUC 一点没涨。训练分一路上行、验证分早早走平，"
              "说明后面的特征更多在帮模型“背训练集”。保留 38 个不增加线上成本；"
              "若将来要精简采集口径，砍到 20–25 个、AUC 损失不到 0.005。"],
         image="feature_count_auc.png"),
    dict(tag="概率校准（W4）",
         title="概率校准：让「概率」能直接算钱",
         lead="模型排序很准，但原始概率整体虚高——平均说 45%，实际只有 20%。"
              "校准后“说多少就真发生多少”。",
         table=[["版本", "平均预测概率", "Brier（越小越好）"],
                ["校准前（原始概率）", "0.4495", "0.2106"],
                ["校准后（Platt）", "0.1993", "0.1424"],
                ["测试集实际违约率", "0.1995", "—"]],
         box=["为什么非要校准",
              "① 概率虚高 125%：直接拿去算期望损失会高估损失，把该放的好客户也拒掉；"
              "② 阈值失真：0.5 本该是“一半一半”的分界，却拒掉了 41% 的申请；"
              "③ AUC 一点没变（0.7249）——校准只动数值大小，不动谁排前面。"
              "方法在验证集内部对半选出：Isotonic 与 Platt 差万分之几，取更平滑的 Platt。"],
         image="calibration_curve.png"),
    dict(tag="期望损失框架（W4）",
         title="从“拍一条线”到“按笔算账”",
         lead="每 1 元贷款的期望利润 = (1−p) × 净收益率 × 期限 − p × (1−回收率)："
              "金额大、利率高、期限长的贷款赚得多，能容忍更高的违约概率。",
         table=[["策略", "被拒占比", "总期望利润（万元）"],
                ["全放（不筛选）", "0%", "2,085"],
                ["固定阈值 0.5", "41.4%", "6,155"],
                ["期望损失框架（中性）", "35.2%", "6,611"]],
         box=["三条读数",
              "① 校准是前提：用未校准概率去算账会拒掉 87.7% 的申请；"
              "② 中性口径（净息差 6%、回收率 30%）下拒绝率 35.2%——比固定阈值少拒 6 个百分点，"
              "期望利润反而从 6,155 万升到 6,611 万，少拒还多赚；"
              "③ 口径比回收率敏感得多（净息差 3%→9%，拒绝率在 65% 与 18% 之间摆动），"
              "先和财务敲定口径，再谈阈值。"],
         image="expected_loss_sensitivity.png"),
]

new_indices = []
for spec in CONTENT:
    s = clone_slide(template)
    new_indices.append(len(prs.slides) - 1)
    # 标签 / 标题 / 导语
    find_shape(s, "10 ·").text_frame.paragraphs[0].runs[0].text = spec["tag"]
    set_lines(find_shape(s, "没有过拟合"), [spec["title"]])
    set_lines(find_shape(s, "模型训练时从没见过"), [spec["lead"]])
    # 表格
    tbl = next(sh for sh in s.shapes if sh.has_table).table
    for r, row in enumerate(spec["table"]):
        for c, val in enumerate(row):
            set_cell(tbl.cell(r, c), val)
    # 要点框（原第 11 页的"三线紧贴说明什么"）
    set_lines(find_shape(s, "三线紧贴说明什么"), spec["box"])
    # 换图
    pic = next(sh for sh in s.shapes if sh.shape_type == 13)
    replace_picture(s, pic, FIG / spec["image"])
    print(f"  新页已生成: {spec['tag']}")

# ---------- 2) 重排为最终顺序 ----------
base_ids = list(range(len(prs.slides) - NEW_SLIDES))       # 0..14 原始页
n1, n2, n3 = new_indices                                  # 15,16,17
order = base_ids[:7] + [n1] + base_ids[7:12] + [n2] + base_ids[12:13] + [n3] + base_ids[13:]
reorder(order)
print(f"已重排为 {len(order)} 页")

# ---------- 3) 同步已有页上过时的数字 ----------
overview = prs.slides[1]
set_lines(find_shape(overview, "模型怎么用？"),
          ["模型怎么用？",
           "概率先校准（说 20% 就真违约 20%），再按每笔贷款的期望利润定审批线："
           "中性口径下拒绝率 35%、比一刀切少拒还多赚，落地走分层审批。"])

threshold = prs.slides[14]
set_lines(find_shape(threshold, "金额账（演示测算"),
          ["金额账（演示测算，非真实业务数据）",
           "口径：坏账损失 = 概率（已校准）× 金额 × (1−回收率)；净收益 = 利息收入 − 坏账损失（详见下页）。",
           "演示参数：回收率 30%、年净息差 6%（机构填真实值后结论即生效）。",
           "· 固定阈值 0.5 → 拒绝 41.4% 申请，期望利润 6,155 万",
           "· 期望损失框架 → 拒绝 35.2% 申请，期望利润 6,611 万（少拒还多赚）",
           "· 口径是最大变量：净息差 3%→9%，拒绝率在 65% 与 18% 之间摆动",
           "落地：分层审批——低分直接拒、中分人工、高分自动通过，模型当排序工具而非最终裁判。"])

advice = prs.slides[16]
set_lines(find_shape(advice, "引入线性"), [
    "引入线性 / 深度模型做异质集成；把校准迁移到真实客群数据（本项目已用 Platt 对齐测试集口径）；"
    "按发放月份做时间序列回测。"])
limits = find_shape(advice, "模型局限（讲给业务听）")
set_lines(limits, [
    "模型局限（讲给业务听）",
    "· 违约率 20% 高于真实客群（通常 <5%）→ 已做 Platt 校准（虚高 125% → 均值 0.199 ≈ 实际），"
    "迁移到真实客群仍需重训重校准",
    "· 匿名特征 n0–n14 有信号但无法向监管解释",
    "· 一期截面数据，无法证明未来时点仍有效",
    "· testA 无标签，线上表现仍需真实放款回标验证",
    "· 精确率有限（0.5 阈值下误伤较多）→ 用分层审批消化，而非单纯压阈值"])

# ---------- 4) 统一重排页码与章节号 ----------
for i, s in enumerate(prs.slides, start=1):
    for sh in text_shapes(s):
        txt = sh.text_frame.text.strip()
        if re.fullmatch(r"\d+\s*/\s*\d+", txt):
            set_para_text(sh.text_frame.paragraphs[0], f"{i} / {TOTAL}")
        elif re.match(r"^\d{2} · ", txt):
            title = txt.split("·", 1)[1].strip()
            set_para_text(sh.text_frame.paragraphs[0], f"{i - 1:02d} · {title}")

prs.save(str(PPTX))
print(f"已保存: {PPTX}（{len(prs.slides)} 页）")
